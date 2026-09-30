"""Dung file cong trinh / mau cong trinh bang openpyxl, dung dung hop dong va cong thuc chung voi VBA.

Day la cong cu phia may phat trien de: (1) tao file DEMO khong can Excel, (2) kiem chung cong thuc
bang LibreOffice. Logic chen cong tac o day phai giong modEstimate.bas.
"""
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import engine
from .errors import TAError
from .schema import load_formulas, load_schema, load_strings, table_map

FORMATS = {"money": "#,##0", "qty": "#,##0.0000", "amount": "#,##0.0000", "rate": "0.00"}
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(color="0000FF")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
CAPTION_FONT = Font(bold=True, size=9, color="1F3864")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
NOTE_FONT = Font(italic=True, size=9, color="7F7F7F")
DV_LAST_ROW = 3000
WIDE_SHEETS = ("TIEN_LUONG", "CHIET_TINH", "GIA_DAU_VAO", "DM_SNAPSHOT")  # in A3; chang 07 chot theo mau cong ty
GROUP_TOKENS = ("VL", "NC", "M")
SUMMARY_GROUPS = ("VL", "NC", "M")


def num(value):
    """Chuoi thap phan chuan -> so Excel (int neu nguyen)."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    d = engine.D(value)
    return int(d) if d == d.to_integral_value() else float(d)


class Project:
    def __init__(self, now=None, user="TA-build"):
        self.schema = load_schema()
        self.formulas = load_formulas()
        self.strings = load_strings()
        self.tables = table_map(self.schema)
        self.rows = {name: [] for name in self.tables}
        self.user = user
        self.now = now or datetime.now().isoformat(timespec="seconds")
        self.counters = {}
        for pk in self.schema["project_keys"]:
            self.rows["tblProject"].append({"key": pk["key"], "label": pk["label"],
                                            "value": pk["value"], "source_ref": ""})

    # ------------------------------------------------------------ tien ich
    def _next(self, prefix, width):
        self.counters[prefix] = self.counters.get(prefix, 0) + 1
        return "%s%0*d" % (prefix, width, self.counters[prefix])

    def add_row(self, table, values):
        names = self.tables[table]["names"]
        unknown = set(values) - set(names)
        if unknown:
            raise KeyError("Cot khong co trong %s: %s" % (table, sorted(unknown)))
        row = {c: values.get(c) for c in names}
        self.rows[table].append(row)
        return row

    def get_project(self, key):
        for r in self.rows["tblProject"]:
            if r["key"] == key:
                return r["value"]
        raise KeyError(key)

    def set_project(self, key, value):
        for r in self.rows["tblProject"]:
            if r["key"] == key:
                r["value"] = value
                return
        raise KeyError(key)

    def audit(self, action, item_id="", field="", before="", after="", reason="", source_ref=""):
        self.add_row("tblAudit", {"event_id": self._next("EV", 6), "timestamp": self.now, "user": self.user,
                                  "action": action, "item_id": item_id, "field": field, "before": before,
                                  "after": after, "reason": reason, "source_ref": source_ref})

    # ------------------------------------------------------------ chen cong tac
    def insert_item(self, catalog, norm_id, input_quantity=None, input_unit=None, unit_conversion="1",
                    conversion_basis="", group_id="", allow_review=False, description=None):
        norms = {n["norm_id"]: n for n in catalog["norms"]}
        if norm_id not in norms:
            raise TAError("NORM_NOT_FOUND", norm_id)
        norm = norms[norm_id]
        if norm["status"] != "verified" and not allow_review:
            raise TAError("NORM_NOT_VERIFIED", "%s dang %s" % (norm["code"], norm["status"]))
        pinned = self.get_project("catalog_release_id")
        if not pinned:
            self.set_project("catalog_release_id", catalog["catalog_release_id"])
        elif pinned != catalog["catalog_release_id"]:
            raise TAError("CATALOG_MISMATCH", "%s <> %s" % (pinned, catalog["catalog_release_id"]))
        resources = {r["resource_id"]: r for r in catalog["resources"]}

        sid = self._next("SN", 4)
        self.add_row("tblSnapshotNorms", {
            "snapshot_id": sid, "norm_id": norm["norm_id"], "catalog_release_id": catalog["catalog_release_id"],
            "namespace": norm["namespace"], "document_id": norm["document_id"],
            "catalog_revision": norm["revision"], "code": norm["code"], "description": norm["name"],
            "source_unit_text": norm["source_unit_text"], "base_unit": norm["base_unit"],
            "norm_basis_qty": num(norm["norm_basis_qty"]), "vl_status": norm["vl_status"],
            "nc_status": norm["nc_status"], "m_status": norm["m_status"], "conditions": norm.get("conditions", ""),
            "source_ref": norm.get("source_ref", ""), "status": norm["status"],
            "is_demo": bool(norm.get("is_demo")), "pinned_at": self.now})

        def global_basis(refs):
            if not refs:
                return ""
            return ";".join(r if r in GROUP_TOKENS else "%s.%s" % (sid, r) for r in refs)

        for c in norm["consumptions"]:
            res = resources.get(c.get("resource_id")) or {}
            self.add_row("tblSnapshotLines", {
                "snapshot_line_id": "%s.%s" % (sid, c["line_id"]), "snapshot_id": sid,
                "local_line_id": c["line_id"], "resource_id": c.get("resource_id") or "",
                "resource_name": res.get("name", "") if c["amount_kind"] == "quantity" else self._percent_label(c),
                "resource_unit": res.get("unit", "") if c["amount_kind"] == "quantity" else "%",
                "group": c["group"], "amount_kind": c["amount_kind"], "amount": num(c["amount"]),
                "amount_text": str(c["amount"]), "basis_ref": global_basis(c.get("basis_ref")),
                "status": c.get("status", norm["status"]), "source_ref": c.get("source_ref", "")})

        scope_order = {"document": 0, "chapter": 1, "group": 2, "norm": 3}
        keys = {"document": norm["document_id"], "chapter": norm.get("chapter"),
                "group": norm.get("group_name"), "norm": norm["norm_id"]}
        notes = [n for n in catalog.get("notes", []) if keys.get(n["scope"]) == n["scope_key"]]
        for n in sorted(notes, key=lambda n: (scope_order[n["scope"]], n["order"])):
            self.add_row("tblSnapshotNotes", {"note_id": "%s.%s" % (sid, n["note_id"]), "snapshot_id": sid,
                                              "scope": n["scope"], "text": n["text"],
                                              "source_ref": n.get("source_ref", "")})

        item_id = self._next("IT", 4)
        self.add_row("tblItems", {
            "item_id": item_id, "group_id": group_id, "code": norm["code"],
            "description": description or norm["name"], "input_unit": input_unit or norm["base_unit"],
            "input_quantity": num(input_quantity), "unit_conversion": num(unit_conversion),
            "conversion_basis": conversion_basis, "quantity_note": "", "norm_id": norm["norm_id"],
            "snapshot_id": sid})

        for c in norm["consumptions"]:
            res = resources.get(c.get("resource_id")) or {}
            is_qty = c["amount_kind"] == "quantity"
            price_id = "P." + c["resource_id"] if is_qty else ""
            self.add_row("tblAnalysis", {
                "analysis_id": self._next("AN", 6), "item_id": item_id,
                "snapshot_line_id": "%s.%s" % (sid, c["line_id"]), "resource_id": c.get("resource_id") or "",
                "resource_spec": res.get("name", "") if is_qty else self._percent_label(c),
                "group": c["group"], "amount_kind": c["amount_kind"],
                "resource_unit": res.get("unit", "") if is_qty else "%", "adjustment_factor": 1,
                "price_id": price_id, "basis_ref": global_basis(c.get("basis_ref")),
                "status": c.get("status", norm["status"])})
            if is_qty:
                self._ensure_price(price_id, res)
        self.audit("insert_item", item_id, "norm_id", "", norm["norm_id"], "",
                   self.strings["audit.insert_source"].replace("{0}", norm["document_id"]).replace("{1}", norm["revision"]))
        return item_id

    def _percent_label(self, c):
        refs = c.get("basis_ref") or []
        return self.strings["label.percent"].replace("{0}", str(c["amount"])).replace("{1}", ";".join(refs))

    def _ensure_price(self, price_id, res):
        if any(p["price_id"] == price_id for p in self.rows["tblPrices"]):
            return
        self.add_row("tblPrices", {"price_id": price_id, "resource_id": res.get("resource_id", ""),
                                   "group": res.get("group", ""), "spec": res.get("name", ""),
                                   "unit": res.get("unit", "")})

    # ------------------------------------------------------------ nhap lieu
    def set_price(self, price_id, **fields):
        for p in self.rows["tblPrices"]:
            if p["price_id"] == price_id:
                for k, v in fields.items():
                    if k not in p:
                        raise KeyError(k)
                    p[k] = num(v) if k in ("quoted_price", "transport", "handling") else v
                return p
        raise KeyError(price_id)

    def apply_factor(self, item_id, group, factor, reason, source_ref=""):
        if group not in GROUP_TOKENS:
            raise TAError("INVALID_GROUP", group)
        if not reason:
            raise TAError("REASON_REQUIRED", "Phai co ly do")
        n = 0
        for a in self.rows["tblAnalysis"]:
            if a["item_id"] == item_id and a["group"] == group and a["amount_kind"] == "quantity":
                before = a["adjustment_factor"]
                a["adjustment_factor"] = num(factor)
                self.audit("apply_factor", item_id, "adjustment_factor:" + a["snapshot_line_id"],
                           before, a["adjustment_factor"], reason, source_ref)
                n += 1
        return n

    def set_cost_rules(self, profile_id, rules, status="draft", source_ref="", reviewer=""):
        self.rows["tblCostRules"] = []
        for r in rules:
            self.add_row("tblCostRules", {
                "rule_id": r["rule_id"], "profile_id": profile_id, "label": r["label"],
                "rule_type": r["rule_type"], "base_refs": r.get("base_refs", ""), "rate": num(r.get("rate")),
                "fixed_amount": num(r.get("fixed_amount")), "rounding_stage": r.get("rounding_stage", "round_cost"),
                "source_ref": r.get("source_ref", source_ref), "reviewer": reviewer, "status": status})
        self.set_project("cost_rule_profile_id", profile_id)

    # ------------------------------------------------------------ tong hop
    def refresh_resource_summary(self):
        self.rows["tblResourceSummary"] = []
        seen = set()
        for a in self.rows["tblAnalysis"]:
            if a["amount_kind"] != "quantity":
                continue
            key = (a["resource_id"], a["price_id"])
            if key in seen:
                continue
            seen.add(key)
            self.add_row("tblResourceSummary", {"resource_price_key": "%s|%s" % key, "resource_id": a["resource_id"],
                                                "group": a["group"], "spec": a["resource_spec"],
                                                "unit": a["resource_unit"], "price_id": a["price_id"]})

    def refresh_summary(self):
        """Ghi TONG_HOP: VL, NC, M, T roi cac quy tac theo thu tu topo; tham chieu o truc tiep de QS doc."""
        t = self.tables["tblSummary"]
        col = {name: get_column_letter(t["first_col"] + i) for i, name in enumerate(t["names"])}
        first = t["header_row"] + 1
        s = self.strings
        rows, where = [], {}

        def add(values):
            where[values["cost_id"]] = first + len(rows)
            rows.append(values)

        for g in SUMMARY_GROUPS:
            add({"cost_id": g, "item_group": s["summary.scope_all"], "cost_label": s["summary." + g],
                 "basis": "CHIET_TINH", "rate": None,
                 "amount": self.formulas["summary.group_sum"].replace("{GROUP}", g),
                 "source_ref": s["summary.source_ct"], "status": "auto"})
        add({"cost_id": "T", "item_group": s["summary.scope_all"], "cost_label": s["summary.T"],
             "basis": "VL;NC;M", "rate": None,
             "amount": "=" + "+".join("%s%d" % (col["amount"], where[g]) for g in SUMMARY_GROUPS),
             "source_ref": "", "status": "auto"})
        rules = [r for r in self.rows["tblCostRules"] if r["rule_id"]]
        ordered = engine.order_cost_rules([{k: (str(v) if v is not None and k in ("rate", "fixed_amount") else v)
                                            for k, v in r.items()} for r in rules])
        for r in ordered:
            refs = [x for x in (r["base_refs"] or "").split(";") if x]
            base = "+".join("%s%d" % (col["amount"], where[x]) for x in refs)
            rid = r["rule_id"]
            row = first + len(rows)
            if r["rule_type"] == "percentage_on_basis":
                rate = self.formulas["summary.rule_rate"].replace("{RULE}", rid)
                expr = "%s%d/100*(%s)" % (col["rate"], row, base)
                amount = "=ROUND(%s,TA_RoundCost)" % expr if r["rounding_stage"] == "round_cost" else "=" + expr
            elif r["rule_type"] == "sum":
                rate, amount = None, "=" + base
            else:
                rate, amount = None, self.formulas["summary.rule_fixed"].replace("{RULE}", rid)
            add({"cost_id": rid, "item_group": s["summary.scope_all"], "cost_label": r["label"],
                 "basis": r["base_refs"], "rate": rate, "amount": amount,
                 "source_ref": r["source_ref"], "status": r["status"]})
        self.rows["tblSummary"] = [{c: v.get(c) for c in t["names"]} for v in rows]
        return where

    # ------------------------------------------------------------ ghi file
    def save(self, path, template=False):
        wb = Workbook()
        wb.remove(wb.active)
        sheets = {}
        for sh in self.schema["sheets"]:
            ws = wb.create_sheet(sh["name"])
            sheets[sh["name"]] = ws
            ws["A1"] = sh["title"]
            ws["A1"].font = TITLE_FONT
            note = {"DM_SNAPSHOT": "sheet.note.snapshot", "KIEM_TRA": "sheet.note.checks"}.get(sh["name"], "sheet.note.input")
            if sh["name"] != "TONG_HOP":
                ws["A2"] = self.strings[note]
                ws["A2"].font = NOTE_FONT
        for name, t in self.tables.items():
            self._write_table(sheets[t["sheet"]], name, t)
        self._write_banner(sheets["TONG_HOP"])
        for dn, expr in self.schema["defined_names"].items():
            wb.defined_names[dn] = DefinedName(dn, attr_text=expr)
        self._validations(sheets)
        for sh in self.schema["sheets"]:
            ws = sheets[sh["name"]]
            hr = sh["header_row"]
            ws.freeze_panes = ws.cell(row=hr + 1, column=2)
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = 8 if sh["name"] in WIDE_SHEETS else 9  # A3 / A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = "%d:%d" % (hr, hr)
            ws.oddFooter.center.text = "Trang &P/&N"
            ws.oddFooter.left.text = "TA Estimate – &A"
        wb.template = template
        wb.calculation.fullCalcOnLoad = True
        wb.save(path)
        return path

    def _write_table(self, ws, name, t):
        hr, c0 = t["header_row"], t["first_col"]
        cols = t["columns"]
        for i, c in enumerate(cols):
            col = c0 + i
            cap = ws.cell(row=hr - 1, column=col, value=c["caption"])
            cap.font = CAPTION_FONT
            cap.alignment = Alignment(wrap_text=True, vertical="bottom")
            h = ws.cell(row=hr, column=col, value=c["name"])
            h.font = Font(bold=True, size=9)
            h.fill = HEADER_FILL
            ws.column_dimensions[get_column_letter(col)].width = c.get("width", 12)
        ws.row_dimensions[hr - 1].height = 30
        rows = self.rows[name]
        for r_i, row in enumerate(rows):
            r = hr + 1 + r_i
            for i, c in enumerate(cols):
                cell = ws.cell(row=r, column=c0 + i)
                val = row.get(c["name"])
                if c["kind"] == "formula" and val is None and name != "tblSummary":
                    val = self.formulas["%s.%s" % (name, c["name"])]
                cell.value = val
                self._style(cell, c)
        n = max(len(rows), 1)
        if not rows:
            for i, c in enumerate(cols):
                self._style(ws.cell(row=hr + 1, column=c0 + i), c)
        ref = "%s%d:%s%d" % (get_column_letter(c0), hr, get_column_letter(c0 + len(cols) - 1), hr + n)
        tab = Table(displayName=name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(tab)

    @staticmethod
    def _style(cell, c):
        if c.get("format") in FORMATS:
            cell.number_format = FORMATS[c["format"]]
        if c["kind"] == "input":
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
            cell.protection = Protection(locked=False)

    def _write_banner(self, ws):
        s, f = self.strings, self.formulas
        ws["A2"], ws["F2"] = s["banner.project"], f["banner.project_name"]
        ws["A3"], ws["F3"] = s["banner.demo"], f["banner.demo_norms"]
        ws["A4"], ws["F4"] = s["banner.errors"], f["banner.error_lines"]
        ws["A5"], ws["F5"] = s["banner.draft"], f["banner.draft_direct"]
        ws["F5"].number_format = FORMATS["money"]
        ws["A6"] = '="%s"&IF(F3>0,"%s",IF(F4>0,"%s","%s"))' % (
            s["banner.status_prefix"], s["banner.status_demo"], s["banner.status_error"], s["banner.status_ok"])
        ws["A6"].font = Font(bold=True, color="C00000")
        for r in range(2, 6):
            ws.cell(row=r, column=1).font = Font(bold=True, size=10)
            ws.cell(row=r, column=6).font = Font(bold=True, size=10)

    def _validations(self, sheets):
        lists = self.schema["lists"]
        spec = [("tblPrices", "tax_basis", "tax_basis"), ("tblPrices", "includes_transport", "includes_transport"),
                ("tblCostRules", "rule_type", "rule_type"), ("tblCostRules", "rounding_stage", "rounding_stage"),
                ("tblCostRules", "status", "rule_status")]
        for table, column, lst in spec:
            t = self.tables[table]
            ws = sheets[t["sheet"]]
            letter = get_column_letter(t["first_col"] + t["names"].index(column))
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(lists[lst]), allow_blank=True)
            dv.add("%s%d:%s%d" % (letter, t["header_row"] + 1, letter, DV_LAST_ROW))
            ws.add_data_validation(dv)
