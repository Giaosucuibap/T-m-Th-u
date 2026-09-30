"""Ghi gia tri da tinh (tu LibreOffice) vao cac o cong thuc cua file openpyxl.

openpyxl chi ghi cong thuc, khong co gia tri: Excel o che do Protected View va trinh xem tren
dien thoai se thay o trong. File van giu fullCalcOnLoad=1 nen Excel tinh lai khi mo chinh sua.
"""
import posixpath
import zipfile

from lxml import etree
from openpyxl import load_workbook

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN}
EXCEL_ERRORS = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NUM!", "#NAME?", "#NULL!"}


def _sheet_parts(z):
    wb = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall("{%s}Relationship" % PKG_REL)}
    out = {}
    for s in wb.find("m:sheets", NS):
        t = target[s.get("{%s}id" % REL)]
        out[posixpath.normpath(posixpath.join("xl", t)) if not t.startswith("/") else t[1:]] = s.get("name")
    return out


def inject(src_path, values_path, out_path):
    values = load_workbook(values_path, data_only=True)
    n = 0
    with zipfile.ZipFile(src_path) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        parts = _sheet_parts(zin)
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in parts:
                ws = values[parts[info.filename]]
                root = etree.fromstring(data)
                for c in root.iterfind(".//m:c", NS):
                    if c.find("m:f", NS) is None:
                        continue
                    v = ws[c.get("r")].value
                    old = c.find("m:v", NS)
                    if old is not None:
                        c.remove(old)
                    if v is None:
                        continue
                    node = etree.SubElement(c, "{%s}v" % MAIN)
                    if isinstance(v, bool):
                        c.set("t", "b")
                        node.text = "1" if v else "0"
                    elif isinstance(v, (int, float)):
                        c.attrib.pop("t", None)
                        node.text = repr(float(v)) if isinstance(v, float) else str(v)
                    elif isinstance(v, str) and v in EXCEL_ERRORS:
                        c.set("t", "e")
                        node.text = v
                    else:
                        c.set("t", "str")
                        node.text = str(v)
                    n += 1
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, data)
    return n
