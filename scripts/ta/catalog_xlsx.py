"""Ghi Catalog_<release>.xlsx (chi gia tri) tu du lieu catalog da kiem."""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .project import num
from .schema import ROOT


def load_catalog_contract():
    with open(os.path.join(ROOT, "contracts", "catalog_tables.json"), encoding="utf-8") as f:
        return json.load(f)


def catalog_rows(cat, created_at):
    lines = []
    for n in cat["norms"]:
        for c in n["consumptions"]:
            lines.append({
                "norm_id": n["norm_id"], "line_id": c["line_id"], "resource_id": c.get("resource_id") or "",
                "group": c["group"], "amount_kind": c["amount_kind"], "amount": num(c["amount"]),
                "amount_text": str(c["amount"]), "basis_ref": ";".join(c.get("basis_ref") or []),
                "status": c.get("status", n["status"]), "source_ref": c.get("source_ref", "")})
    norms = []
    for n in cat["norms"]:
        row = {k: n.get(k, "") for k in ("norm_id", "namespace", "document_id", "revision", "code", "variant_id",
                                         "appendix", "chapter", "group_name", "name", "base_unit",
                                         "source_unit_text", "status", "vl_status", "nc_status", "m_status",
                                         "conditions", "source_ref")}
        row["norm_basis_qty"] = num(n["norm_basis_qty"])
        row["is_demo"] = bool(n.get("is_demo"))
        norms.append(row)
    manifest = [
        {"key": "catalog_release_id", "value": cat["catalog_release_id"]},
        {"key": "catalog_schema_version", "value": "1.0"},
        {"key": "is_demo", "value": bool(cat.get("is_demo"))},
        {"key": "created_at", "value": created_at},
        {"key": "warning", "value": cat.get("warning", "")},
        {"key": "norm_count", "value": len(norms)},
        {"key": "line_count", "value": len(lines)},
        {"key": "supported_codes_note", "value": "Chỉ các mã trong bảng NORMS. Không phải thư viện đầy đủ."},
    ]
    return {"tblCatManifest": manifest, "tblCatDocuments": cat["documents"], "tblCatNorms": norms,
            "tblCatResources": cat["resources"], "tblCatLines": lines, "tblCatNotes": cat.get("notes", [])}


def write_catalog(cat, path, created_at):
    contract = load_catalog_contract()
    data = catalog_rows(cat, created_at)
    wb = Workbook()
    wb.remove(wb.active)
    for t in contract["tables"]:
        ws = wb.create_sheet(t["sheet"])
        cols = t["columns"]
        for i, c in enumerate(cols, start=1):
            ws.cell(row=1, column=i, value=c).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(i)].width = 14 if c not in ("name", "text", "conditions") else 50
        rows = data[t["name"]]
        for r, row in enumerate(rows, start=2):
            for i, c in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=i, value=row.get(c))
                if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
                    cell.data_type = "s"  # du lieu nguon luon la chu, khong bao gio thanh cong thuc
        ref = "A1:%s%d" % (get_column_letter(len(cols)), max(len(rows), 1) + 1)
        tab = Table(displayName=t["name"], ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
        ws.freeze_panes = "A2"
    wb.save(path)
    return path
