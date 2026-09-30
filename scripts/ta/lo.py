"""Tinh lai workbook bang LibreOffice headless de kiem chung cong thuc (khong thay the Excel that)."""
import os
import shutil
import subprocess
import tempfile

from openpyxl import load_workbook


def soffice_path():
    return shutil.which("soffice") or shutil.which("libreoffice")


def available():
    return soffice_path() is not None


def recalc(path, outdir=None):
    """Mo bang LibreOffice, tinh lai, luu .xlsx moi. Tra ve duong dan file ket qua."""
    outdir = os.path.abspath(outdir or tempfile.mkdtemp(prefix="ta_lo_"))
    os.makedirs(outdir, exist_ok=True)
    # Ho so LibreOffice dung chung, tach khoi ho so nguoi dung; lan dau khoi tao doi khi loi -> thu lai.
    profile = os.path.join(tempfile.gettempdir(), "ta_lo_profile")
    os.makedirs(profile, exist_ok=True)
    env = dict(os.environ, HOME=profile)
    src = os.path.abspath(path)
    base = os.path.splitext(os.path.basename(src))[0] + ".xlsx"
    out = os.path.join(outdir, base)
    if os.path.abspath(out) == src:
        raise ValueError("outdir trung thu muc file goc")
    cmd = [soffice_path(), "-env:UserInstallation=file://" + profile, "--headless", "--norestore",
           "--nologo", "--convert-to", "xlsx", "--outdir", outdir, src]
    log = ""
    for _ in range(3):
        proc = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=300)
        if os.path.exists(out):
            return out
        log = proc.stdout + proc.stderr
    raise RuntimeError("LibreOffice khong tao duoc file: %s" % log[-2000:])


def recalc_many(paths, outdir):
    """Tinh lai nhieu file trong mot lan goi LibreOffice. Tra ve list duong dan ket qua."""
    outdir = os.path.abspath(outdir)
    os.makedirs(outdir, exist_ok=True)
    profile = os.path.join(tempfile.gettempdir(), "ta_lo_profile")
    os.makedirs(profile, exist_ok=True)
    env = dict(os.environ, HOME=profile)
    srcs = [os.path.abspath(p) for p in paths]
    outs = [os.path.join(outdir, os.path.splitext(os.path.basename(s))[0] + ".xlsx") for s in srcs]
    cmd = [soffice_path(), "-env:UserInstallation=file://" + profile, "--headless", "--norestore",
           "--nologo", "--convert-to", "xlsx", "--outdir", outdir] + srcs
    log = ""
    for _ in range(3):
        proc = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=600)
        if all(os.path.exists(o) for o in outs):
            return outs
        log = proc.stdout + proc.stderr
    raise RuntimeError("LibreOffice khong tao du file: %s" % log[-2000:])


def open_values(path):
    return load_workbook(path, data_only=True)


def read_table(src, sheet, header_row, first_col, names):
    """Doc gia tri bang theo vi tri header trong file da tinh lai. Dung khi gap o trong o cot khoa.

    src: duong dan hoac workbook da mo bang open_values().
    """
    wb = open_values(src) if isinstance(src, str) else src
    ws = wb[sheet]
    header = [ws.cell(row=header_row, column=first_col + i).value for i in range(len(names))]
    if header != names:
        raise AssertionError("Header %s khong khop hop dong: %s" % (sheet, header))
    rows = []
    r = header_row + 1
    while True:
        vals = [ws.cell(row=r, column=first_col + i).value for i in range(len(names))]
        if vals[0] in (None, ""):
            break
        rows.append(dict(zip(names, vals)))
        r += 1
    return rows


def read_cells(src, sheet, refs):
    wb = open_values(src) if isinstance(src, str) else src
    ws = wb[sheet]
    return {ref: ws[ref].value for ref in refs}
