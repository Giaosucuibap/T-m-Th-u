"""Independent OOXML reader verification of synthetic production-export fixtures."""
import hashlib
import json
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
directory = root / 'test-results' / 'excel-multi-investor-414'
source = json.loads((directory / 'export-fixture-verification.json').read_text(encoding='utf-8'))
checks = []
for entry in source['files']:
    path = directory / entry['filename']
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    assert before == entry['sha256'], entry['filename']
    book = openpyxl.load_workbook(path, read_only=False, data_only=False)
    detail = book[entry['detailSheet']]
    assert detail.max_row == 4, (entry['kind'], detail.max_row)
    rows = list(detail.values)
    owner = book['Đối soát']
    values = {row[0]: row[1] for row in owner.values if row[0]}
    assert values['Tiêu chí gốc · Chủ đầu tư'] == source['criteria']['investor']
    assert values['Tiêu chí gốc · Tỉnh/thành phố'] == source['criteria']['province']
    assert [values[name] for name in ['e-GP công bố', 'Đã tải từ e-GP', 'Khớp tiêu chí gốc', 'Chưa đủ dữ liệu', 'Ngoài tiêu chí gốc']] == [5, 5, 3, 1, 1]
    assert values['Số bản ghi sau bộ lọc xuất'] == 3
    assert values['Số dòng dữ liệu xuất'] == 3
    assert detail.freeze_panes and detail.auto_filter.ref
    for row in detail.iter_rows(min_row=2):
        assert any(cell.border.bottom.style == 'thin' for cell in row)
        assert not any(cell.data_type == 'f' for cell in row)
    header = {cell.value: cell.column for cell in detail[1]}
    money_header = next(name for name in ['Giá gói thầu', 'Giá gói thầu/dự toán', 'Giá trúng thầu'] if name in header)
    assert all(detail.cell(row, header[money_header]).data_type == 'n' for row in range(2, 5))
    identifiers = []
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != 'f'
                if cell.value == '0012345678':
                    assert cell.data_type == 's'
                    identifiers.append(f'{sheet.title}!{cell.coordinate}')
    if entry['kind'] in ['BBMT', 'CHU-DAU-TU']:
        assert identifiers
    if entry['kind'] in ['KQLCNT', 'DIA-BAN', 'CHU-DAU-TU']:
        unknown = book['Chưa đủ dữ liệu']
        assert unknown.max_row == 2
        texts = '\n'.join(str(cell.value or '') for row in unknown for cell in row)
        assert 'IB2600000104-00' in texts
        assert 'IB2600000105-00' not in texts
        assert all('IB2600000104' not in str(cell or '') for row in rows for cell in row)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    checks.append({'kind': entry['kind'], 'sha256': before, 'sheets': len(book.sheetnames),
                   'matchedRows': 3, 'unknownRows': 1 if 'Chưa đủ dữ liệu' in book.sheetnames else 0,
                   'leadingZeroCells': identifiers, 'passed': True})
    book.close()

report = {'ok': True, 'reader': 'openpyxl', 'synthetic': True, 'files': checks}
(directory / 'independent-reader-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'ok': True, 'workbooks': len(checks), 'sheets': sum(item['sheets'] for item in checks)}))
